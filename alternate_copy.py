
from openpyxl import load_workbook

# Open the first Excel file
wb1 = load_workbook("Book1.xlsx")
s1 = wb1.active

# Open the second Excel file
wb2 = load_workbook("Book2.xlsx")
s2 = wb2.active

# Copy alternate rows
r = 1

for i in range(1, s1.max_row + 1, 2):
    for j in range(1, s1.max_column + 1):
        s2.cell(r, j).value = s1.cell(i, j).value
    r = r + 1

# Save the second Excel file
wb2.save("Book2.xlsx")

print("Alternate rows copied successfully")